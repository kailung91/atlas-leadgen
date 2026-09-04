"""
Classify 1C UTP and 1C Magazin Counterparties by Order Type & Purchasing Behavior.

Categories:
  1. 🎨 Карти під замовлення (Custom wall maps, lamination, mounting, services)
  2. 🔄 Регулярні замовлення (Regular recurring buyers: ops >= 5 and standard product lines)
  3. ⚡ Разові замовлення (One-off / single purchase buyers: ops < 5)
"""

from __future__ import annotations

import re
from pathlib import Path
import pandas as pd
from loguru import logger
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT / "output"


def norm_name(s: str | None) -> str:
    if pd.isna(s) or not s:
        return ""
    return re.sub(r"\s+", " ", str(s).strip().lower())


def classify_counterparty(row: dict) -> tuple[str, str, str]:
    nom = str(row.get("Основна номенклатурна група", "") or "").lower()
    group_1c = str(row.get("Група Контрагента", "") or "").lower()
    ops = row.get("Кількість операцій", 0)
    if pd.isna(ops):
        ops = 0

    custom_keywords = [
        "ламінація", "планки", "заказ", "послуги", "роздруковка",
        "спец", "друк", "індивід", "монтаж", "оформлення"
    ]

    is_custom = any(kw in nom for kw in custom_keywords) or any(kw in group_1c for kw in custom_keywords)

    if is_custom:
        cat = "🎨 Карти під замовлення (Спецзамовлення/Ламінація)"
        order_type = "Спецзамовлення / Послуга"
        rec_offer = "Каталог великоформатних стінних карт, ламінації та планок під спецзамовлення"
    elif ops >= 5:
        cat = "🔄 Регулярні замовлення"
        order_type = "Гурт / Серійні покупки"
        rec_offer = "Каталог нових атласів/карт + гуртові знижки на новий навчальний рік"
    else:
        cat = "⚡ Разові замовлення"
        order_type = "Разовий Покупець"
        rec_offer = "Реактивація: безкоштовна доставка та спецціна на мінімальну партію"

    return cat, order_type, rec_offer


def format_excel_sheet(ws, df: pd.DataFrame, header_color: str = "1F4E78"):
    ws.freeze_panes = "A2"
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    for col_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, col_name in enumerate(df.columns, 1):
        sample_vals = df[col_name].dropna().astype(str).head(100)
        max_len = max([len(str(v)) for v in sample_vals] + [len(col_name)])
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)


def process_1c_dataset(base_file: Path, verified_file: Path, out_excel: Path, out_csv: Path, db_label: str, header_color: str, name_col_label: str):
    if not base_file.exists() or not verified_file.exists():
        logger.warning("Missing required files for {}: {} or {}", db_label, base_file, verified_file)
        return None

    logger.info("Processing & Classifying 1C [{}] Counterparties...", db_label)
    df_base = pd.read_excel(base_file)
    df_ver = pd.read_excel(verified_file)

    df_base["_norm_name"] = df_base[name_col_label].apply(norm_name)
    df_ver["_norm_name"] = df_ver[name_col_label].apply(norm_name)

    # Verified columns to pull in
    ver_cols = ["_norm_name", "ЄДРПОУ", "ІПН", "Телефони", "Email", "Адреси", "Джерело Email", "Тип Метчингу"]
    ver_cols_present = [c for c in ver_cols if c in df_ver.columns]

    # Deduplicate verified right side by _norm_name
    df_ver_dedup = df_ver.drop_duplicates(subset=["_norm_name"])[ver_cols_present]

    df_merged = pd.merge(df_base, df_ver_dedup, on="_norm_name", how="left")
    df_merged.drop(columns=["_norm_name"], errors="ignore", inplace=True)

    # Apply classification
    categories = []
    order_types = []
    offers = []

    for _, r in df_merged.iterrows():
        cat, ot, off = classify_counterparty(r)
        categories.append(cat)
        order_types.append(ot)
        offers.append(off)

    df_merged["Категорія Клієнта"] = categories
    df_merged["Тип Замовлень"] = order_types
    df_merged["Рекомендований Оффер"] = offers

    # Save output
    df_merged.to_csv(out_csv, index=False, encoding="utf-8-sig")

    with pd.ExcelWriter(out_excel, engine="openpyxl") as writer:
        df_merged.to_excel(writer, index=False, sheet_name=f"1С {db_label} Класифікація")
        format_excel_sheet(writer.sheets[f"1С {db_label} Класифікація"], df_merged, header_color=header_color)

    logger.info("Saved Classified 1C [{}] dataset -> {} and {}", db_label, out_excel.name, out_csv.name)
    return df_merged


def main():
    # 1. Process 1C UTP
    utp_base = OUTPUT_DIR / "1c_base_utp_ipt_contacts.xlsx"
    utp_ver = OUTPUT_DIR / "1c_utp_full_contacts_with_verified_emails.xlsx"
    utp_out_xlsx = OUTPUT_DIR / "1c_utp_full_contacts_classified.xlsx"
    utp_out_csv = OUTPUT_DIR / "1c_utp_full_contacts_classified.csv"

    df_utp = process_1c_dataset(utp_base, utp_ver, utp_out_xlsx, utp_out_csv, "UTP", "1F4E78", "Контрагент (1С UTP)")

    # 2. Process 1C Magazin
    mag_base = OUTPUT_DIR / "1c_base_magazin_contacts.xlsx"
    mag_ver = OUTPUT_DIR / "1c_magazin_full_contacts_with_verified_emails.xlsx"
    mag_out_xlsx = OUTPUT_DIR / "1c_magazin_full_contacts_classified.xlsx"
    mag_out_csv = OUTPUT_DIR / "1c_magazin_full_contacts_classified.csv"

    df_mag = process_1c_dataset(mag_base, mag_ver, mag_out_xlsx, mag_out_csv, "Magazin", "2E75B6", "Контрагент (1С Магазин)")

    # 3. Master Combined Segmented Dataset
    if df_utp is not None and df_mag is not None:
        df_utp["Джерельна База 1С"] = "1С UTP"
        df_mag["Джерельна База 1С"] = "1С Магазин"

        all_cols = list(dict.fromkeys(list(df_utp.columns) + list(df_mag.columns)))
        df_utp_re = df_utp.reindex(columns=all_cols)
        df_mag_re = df_mag.reindex(columns=all_cols)

        df_master = pd.concat([df_utp_re, df_mag_re], ignore_index=True)

        master_xlsx = OUTPUT_DIR / "1c_master_segmented_contacts.xlsx"
        master_csv = OUTPUT_DIR / "1c_master_segmented_contacts.csv"

        df_master.to_csv(master_csv, index=False, encoding="utf-8-sig")

        with pd.ExcelWriter(master_xlsx, engine="openpyxl") as writer:
            df_master.to_excel(writer, index=False, sheet_name="Master 1С Сегментований")
            format_excel_sheet(writer.sheets["Master 1С Сегментований"], df_master, header_color="002060")

        logger.info("Saved Master Combined Segmented Dataset -> {}", master_xlsx.name)


if __name__ == "__main__":
    main()
