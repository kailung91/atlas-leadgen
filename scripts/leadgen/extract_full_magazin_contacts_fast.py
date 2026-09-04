"""
Fast Vectorized 1C Magazin Contact Extraction Pipeline.
"""

from __future__ import annotations

import os
from pathlib import Path
import win32com.client
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from loguru import logger

BASE_DIR = r"G:\Мій диск\Work\База 1С\1C8_Data_Base"
OUT_DIR = Path("D:/Codex/redactor/atlas-industry-engine/output")
OUT_DIR.mkdir(parents=True, exist_ok=True)

QUERY_MAGAZIN_CONTACTS = """
ВЫБРАТЬ
    Контрагенты.Код КАК Код,
    Контрагенты.Наименование КАК Назва,
    Контрагенты.НаименованиеПолное КАК ПовнаНазва,
    Контрагенты.КодПоЕДРПОУ КАК ЄДРПОУ,
    Контрагенты.ИНН КАК ІПН,
    Контрагенты.ОсновноеКонтактноеЛицо.Наименование КАК КонтактнаОсоба,
    КонтактнаяИнформация.Тип КАК Тип,
    КонтактнаяИнформация.Вид.Наименование КАК Вид,
    КонтактнаяИнформация.Представление КАК Представлення
ИЗ
    Справочник.Контрагенты КАК Контрагенты
    ЛЕВОЕ СОЕДИНЕНИЕ РегистрСведений.КонтактнаяИнформация КАК КонтактнаяИнформация
    ПО КонтактнаяИнформация.Объект = Контрагенты.Ссылка
"""


def unwrap(v):
    if v is None:
        return ""
    if isinstance(v, (int, float, bool, str)):
        return str(v).strip()
    return str(v).strip()


def format_excel_sheet(ws, df: pd.DataFrame, header_color: str = "2E75B6") -> None:
    ws.freeze_panes = "A2"
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    for col in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, col_name in enumerate(df.columns, 1):
        sample_vals = df[col_name].dropna().astype(str).head(100)
        max_len = max([len(str(v)) for v in sample_vals] + [len(col_name)])
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)


def extract_magazin_full_contacts():
    conn_str = r'File="G:\Мій диск\Work\База 1С\1C8_Data_Base\Магазин";Usr="admin";Pwd="Cjpthwfybt!33";'
    logger.info("Connecting to 1C Magazin...")

    connector = win32com.client.Dispatch("V83.COMConnector")
    v83 = connector.Connect(conn_str)
    logger.info("Connected to 1C Magazin!")

    q = v83.NewObject("Запрос")
    q.Текст = QUERY_MAGAZIN_CONTACTS
    vt = q.Выполнить().Выгрузить()

    cols = [vt.Колонки.Получить(i).Имя for i in range(vt.Колонки.Количество())]
    total_rows = vt.Количество()
    logger.info("Unloaded ValueTable from 1C Magazin: {} rows!", total_rows)

    rows = []
    for idx in range(total_rows):
        row_obj = vt.Получить(idx)
        rows.append([unwrap(getattr(row_obj, c)) for c in cols])

    del v83, connector
    df = pd.DataFrame(rows, columns=cols)

    # Vectorized classification
    df["Представлення"] = df["Представлення"].fillna("").astype(str).str.strip()
    df["Вид_lower"] = df["Вид"].fillna("").astype(str).str.lower()

    is_email = df["Вид_lower"].str.contains("почта|e-mail|email", regex=True) | df["Представлення"].str.contains("@")
    is_phone = df["Вид_lower"].str.contains("телефон|моб|факс", regex=True) | (df["Представлення"].str.replace(r"[+\-\s\(\)]", "", regex=True).str.isdigit() & (df["Представлення"].str.len() >= 7)) & (~is_email)
    is_addr = df["Вид_lower"].str.contains("адрес|юр|факт", regex=True) & (~is_email) & (~is_phone)

    df["Phone_Val"] = df["Представлення"].where(is_phone, "")
    df["Email_Val"] = df["Представлення"].where(is_email, "")
    df["Addr_Val"] = df["Представлення"].where(is_addr, "")

    def join_unique(s):
        vals = sorted(set(v for v in s if v))
        return "; ".join(vals)

    aggregated = df.groupby(["Код", "Назва"]).agg(
        ПовнаНазва=("ПовнаНазва", "first"),
        ЄДРПОУ=("ЄДРПОУ", "first"),
        ІПН=("ІПН", "first"),
        КонтактнаОсоба=("КонтактнаОсоба", "first"),
        Телефони=("Phone_Val", join_unique),
        Email=("Email_Val", join_unique),
        Адреси=("Addr_Val", join_unique),
    ).reset_index()

    aggregated.rename(columns={
        "Код": "Код 1С",
        "Назва": "Контрагент (1С Магазин)",
        "ПовнаНазва": "Повна Назва",
        "КонтактнаОсоба": "Контактна Особа",
    }, inplace=True)

    logger.info("Consolidated {} unique counterparties for 1C Magazin!", len(aggregated))

    out_csv = OUT_DIR / "1c_magazin_full_contacts.csv"
    aggregated.to_csv(out_csv, index=False, encoding="utf-8-sig")

    out_excel = OUT_DIR / "1c_magazin_full_contacts.xlsx"
    with pd.ExcelWriter(out_excel, engine="openpyxl") as writer:
        aggregated.to_excel(writer, index=False, sheet_name="Контакти 1С Магазин")
        format_excel_sheet(writer.sheets["Контакти 1С Магазин"], aggregated, header_color="2E75B6")

    logger.info("Saved 1C Magazin Full Contacts Excel -> {}", out_excel)
    return aggregated


if __name__ == "__main__":
    extract_magazin_full_contacts()
