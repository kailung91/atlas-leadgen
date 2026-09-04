"""
Comprehensive 1C Contact Information Register Extraction Script.

Extracts:
  - Counterparty Name (Назва) & Full Legal Name (Повна Назва)
  - EDRPOU (Код по ЄДРПОУ) & INN (ІПН)
  - Primary Contact Person (Основне контактне лицо)
  - Telephone Numbers (Телефони)
  - Email addresses (Електронні пошти)
  - Legal / Physical Addresses (Юридичні / Фактичні адреси)
"""

from __future__ import annotations

import os
from pathlib import Path
import win32com.client
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from loguru import logger

BASE_DIR = r"G:\Мій диск\Work\База 1С\1C8_Data_Base"
OUT_DIR = Path("D:/Codex/redactor/atlas-industry-engine/output")
OUT_DIR.mkdir(parents=True, exist_ok=True)

# Correct metadata query for 1C UTP (IPT)
QUERY_UTP_CONTACTS = """
ВЫБРАТЬ
    Контрагенты.Код КАК Код,
    Контрагенты.Наименование КАК Назва,
    Контрагенты.НаименованиеПолное КАК ПовнаНазва,
    Контрагенты.КодПоЕДРПОУ КАК ЄДРПОУ,
    Контрагенты.ИНН КАК ІПН,
    Контрагенты.ОсновноеКонтактноеЛицо.Наименование КАК КонтактнаОсоба,
    КонтактнаяИнформация.Тип КАК Тип,
    КонтактнаяИнформация.Вид.Наименование КАК Вид,
    КонтактнаяИнформация.Представление КАК Представлення
ИЗ
    Справочник.Контрагенты КАК Контрагенты
    ЛЕВОЕ СОЕДИНЕНИЕ РегистрСведений.КонтактнаяИнформация КАК КонтактнаяИнформация
    ПО КонтактнаяИнформация.Объект = Контрагенты.Ссылка
"""


def unwrap(v):
    if v is None:
        return ""
    if isinstance(v, (int, float, bool, str)):
        return str(v).strip()
    return str(v).strip()


def format_excel_sheet(ws, df: pd.DataFrame, header_color: str = "1F4E78") -> None:
    ws.freeze_panes = "A2"
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    for col in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, col_name in enumerate(df.columns, 1):
        sample_vals = df[col_name].dropna().astype(str).head(100)
        max_len = max([len(str(v)) for v in sample_vals] + [len(col_name)])
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)


def extract_utp_full_contacts():
    conn_str = r'File="G:\Мій диск\Work\База 1С\1C8_Data_Base\UTP";Usr="Админ";Pwd="Cjpthwfybt!33";'
    logger.info("Connecting to 1C UTP (IPT)...")

    connector = win32com.client.Dispatch("V83.COMConnector")
    v83 = connector.Connect(conn_str)
    logger.info("Connected to 1C UTP!")

    q = v83.NewObject("Запрос")
    q.Текст = QUERY_UTP_CONTACTS
    executed = q.Выполнить()
    cols = [c.Name for c in executed.Колонки]
    sel = executed.Выбрать()

    rows = []
    while sel.Следующий():
        row = {c: unwrap(getattr(sel, c)) for c in cols}
        rows.append(row)

    del v83, connector
    df_raw = pd.DataFrame(rows)
    logger.info("Extracted {} raw contact records from 1C UTP!", len(df_raw))

    # Pivot / Group by Counterparty so each row is 1 counterparty with combined phones, emails, addresses, EDRPOU
    grouped = []
    for (code, name), group in df_raw.groupby(["Код", "Назва"]):
        legal_name = group["ПовнаНазва"].iloc[0] if not group["ПовнаНазва"].empty else ""
        edrpou = group["ЄДРПОУ"].iloc[0] if not group["ЄДРПОУ"].empty else ""
        inn = group["ІПН"].iloc[0] if not group["ІПН"].empty else ""
        contact_person = group["КонтактнаОсоба"].iloc[0] if not group["КонтактнаОсоба"].empty else ""

        phones = set()
        emails = set()
        addresses = set()

        for _, r in group.iterrows():
            kind = str(r["Вид"]).lower()
            val = str(r["Представлення"]).strip()
            if not val:
                continue

            if "телефон" in kind or "моб" in kind or "факс" in kind or val.replace("+", "").replace("-", "").isdigit():
                phones.add(val)
            elif "почта" in kind or "e-mail" in kind or "email" in kind or "@" in val:
                emails.add(val)
            elif "адрес" in kind or "юр" in kind or "факт" in kind:
                addresses.add(val)

        grouped.append({
            "Код 1С": code,
            "Контрагент (1С UTP)": name,
            "Повна Назва": legal_name,
            "ЄДРПОУ": edrpou,
            "ІПН": inn,
            "Контактна Особа": contact_person,
            "Телефони": "; ".join(sorted(phones)),
            "Email": "; ".join(sorted(emails)),
            "Адреси": "; ".join(sorted(addresses)),
        })

    df_clean = pd.DataFrame(grouped)
    logger.info("Consolidated {} unique counterparties with contacts in 1C UTP!", len(df_clean))

    out_csv = OUT_DIR / "1c_utp_full_contacts.csv"
    df_clean.to_csv(out_csv, index=False, encoding="utf-8-sig")

    out_excel = OUT_DIR / "1c_utp_full_contacts.xlsx"
    with pd.ExcelWriter(out_excel, engine="openpyxl") as writer:
        df_clean.to_excel(writer, index=False, sheet_name="Контакти 1С UTP")
        format_excel_sheet(writer.sheets["Контакти 1С UTP"], df_clean, header_color="1F4E78")

    logger.info("Saved 1C UTP Full Contacts Excel -> {}", out_excel)
    return df_clean


if __name__ == "__main__":
    extract_utp_full_contacts()
