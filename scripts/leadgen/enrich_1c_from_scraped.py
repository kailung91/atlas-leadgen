"""
Enrich 1C UTP (IPT) and 1C Magazin Counterparty Databases with Emails.

Data Sources:
  1. YouControl scraped records (data/raw/youcontrol/*.jsonl)
  2. UA-Region full industrial catalog (data/raw/ua_region_full/*.jsonl)
  3. UA-Region partner leads (data/raw/ua_region_partners/*.jsonl)
  4. Scraped CSV files in Email/ directory
  5. Internal reference databases (Schools, HEIs, Prozorro, Ukrmaps)

Matching Logic:
  1. Primary: ЄДРПОУ (EDRPOU exact match)
  2. Secondary: Назва (Normalized Company Name match)
  3. Tertiary: Телефон (Phone last 9 digits match)
"""

from __future__ import annotations

import glob
import json
import os
import re
import socket
from pathlib import Path
import pandas as pd
from loguru import logger
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT / "output"
RAW_DIR = ROOT / "data" / "raw"
EMAIL_DIR = ROOT / "Email"

HOSTILE_SUFFIXES = ('.ru', '.su', '.by', '.xn--p1ai')
HOSTILE_DOMAINS = ('mail.ru', 'yandex.ru', 'yandex.ua', 'rambler.ru', 'bk.ru', 'inbox.ru', 'list.ru', 'mail.ua')


def is_valid_email_format(email: str) -> bool:
    if not isinstance(email, str) or not email or "@" not in email:
        return False
    email = email.strip().lower()
    if any(email.endswith(sfx) for sfx in HOSTILE_SUFFIXES):
        return False
    domain = email.split("@")[-1]
    if domain in HOSTILE_DOMAINS:
        return False
    return bool(re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email))


def clean_edrpou(raw_edr: str | int | float) -> str:
    if pd.isna(raw_edr) or not raw_edr:
        return ""
    s = str(raw_edr).strip()
    if "." in s:
        s = s.split(".")[0]
    s = s.lstrip("0").strip()
    if len(s) >= 5 and s.isdigit():
        return s
    return ""


def clean_company_name(name: str) -> str:
    if not isinstance(name, str) or not name:
        return ""
    n = name.lower()
    # Strip quotes, brackets, and common legal form prefixes
    n = re.sub(r'[\"\'«»`’\(\)\[\]]', ' ', n)
    n = re.sub(r'\b(тов|пп|фоп|дп|тзов|ао|пао|прат|приватне|товариство|підприємство|фірма|компанія|тдв|кооператив)\b', ' ', n)
    n = re.sub(r'\s+', ' ', n).strip()
    return n


def clean_phone_last9(phone: str) -> str:
    if not isinstance(phone, str) or not phone:
        return ""
    digits = re.sub(r"\D", "", phone)
    if len(digits) >= 9:
        return digits[-9:]
    return ""


def build_master_scraped_email_lookups() -> tuple[dict[str, tuple[str, str]], dict[str, tuple[str, str]], dict[str, tuple[str, str]]]:
    logger.info("Building master scraped email lookup tables from YouControl, UA-Region, and Email datasets...")

    edrpou_map: dict[str, tuple[str, str]] = {}
    name_map: dict[str, tuple[str, str]] = {}
    phone_map: dict[str, tuple[str, str]] = {}

    def add_record(email: str, edrpou_raw: str, name_raw: str, phone_raw: str, source_label: str):
        if not is_valid_email_format(email):
            return
        email = email.strip().lower()

        edr = clean_edrpou(edrpou_raw)
        if edr and edr not in edrpou_map:
            edrpou_map[edr] = (email, source_label)

        cn = clean_company_name(name_raw)
        if cn and len(cn) >= 4 and cn not in name_map:
            name_map[cn] = (email, source_label)

        p9 = clean_phone_last9(phone_raw)
        if p9 and p9 not in phone_map:
            phone_map[p9] = (email, source_label)

    # 1. YouControl JSONL
    yc_files = list((RAW_DIR / "youcontrol").glob("*.jsonl"))
    logger.info("Scanning {} YouControl raw files...", len(yc_files))
    for f in yc_files:
        try:
            with open(f, encoding="utf-8") as file:
                for line in file:
                    if not line.strip():
                        continue
                    rec = json.loads(line)
                    email = rec.get("email") or ""
                    edr = rec.get("edrpou") or rec.get("fop_id") or ""
                    name = rec.get("name") or ""
                    phone = rec.get("phones") or ""
                    add_record(email, str(edr), str(name), str(phone), "YouControl")
        except Exception as e:
            logger.warning("Error reading {}: {}", f.name, e)

    # 2. UA-Region Full Catalog
    uar_full_files = list((RAW_DIR / "ua_region_full").glob("*.jsonl"))
    logger.info("Scanning {} UA-Region Full Catalog raw files...", len(uar_full_files))
    for f in uar_full_files:
        try:
            with open(f, encoding="utf-8") as file:
                for line in file:
                    if not line.strip():
                        continue
                    rec = json.loads(line)
                    email = rec.get("email") or ""
                    edr = rec.get("edrpou") or ""
                    name = rec.get("name") or ""
                    phone = rec.get("phones") or ""
                    add_record(email, str(edr), str(name), str(phone), "UA-Region Full")
        except Exception as e:
            logger.warning("Error reading {}: {}", f.name, e)

    # 3. UA-Region Partners
    uar_partner_files = list((RAW_DIR / "ua_region_partners").glob("*.jsonl"))
    logger.info("Scanning {} UA-Region Partner raw files...", len(uar_partner_files))
    for f in uar_partner_files:
        try:
            with open(f, encoding="utf-8") as file:
                for line in file:
                    if not line.strip():
                        continue
                    rec = json.loads(line)
                    email = rec.get("email") or ""
                    edr = rec.get("edrpou") or ""
                    name = rec.get("name") or ""
                    phone = rec.get("phones") or ""
                    add_record(email, str(edr), str(name), str(phone), "UA-Region Partners")
        except Exception as e:
            logger.warning("Error reading {}: {}", f.name, e)

    # 4. Email Directory CSVs
    if EMAIL_DIR.exists():
        email_csvs = list(EMAIL_DIR.glob("**/*.csv"))
        logger.info("Scanning {} CSV files in Email directory...", len(email_csvs))
        for f in email_csvs:
            try:
                df = pd.read_csv(f, low_memory=False)
                cols = {c.lower(): c for c in df.columns}
                email_col = cols.get("email")
                edr_col = cols.get("edrpou") or cols.get("код єдрпоу") or cols.get("єдрпоу")
                name_col = cols.get("name") or cols.get("назва") or cols.get("компанія")
                phone_col = cols.get("phones") or cols.get("phone") or cols.get("телефон")

                if email_col:
                    for _, r in df.iterrows():
                        em = str(r.get(email_col, "")).strip()
                        ed = str(r.get(edr_col, "")) if edr_col else ""
                        nm = str(r.get(name_col, "")) if name_col else ""
                        ph = str(r.get(phone_col, "")) if phone_col else ""
                        add_record(em, ed, nm, ph, f"Email Catalog ({f.parent.name})")
            except Exception:
                pass

    # 5. Internal Reference CSVs in Output
    for fname in ["school_leads_filtered.csv", "hei_leads_filtered.csv", "prozorro_school_buyers.csv", "prozorro_school_winners.csv", "ukrmaps_ecommerce_customers.csv"]:
        ref_p = OUTPUT_DIR / fname
        if ref_p.exists():
            try:
                df = pd.read_csv(ref_p)
                for _, r in df.iterrows():
                    em = str(r.get("Email", "")).strip()
                    ed = str(r.get("ЄДРПОУ", "")).strip()
                    nm = str(r.get("Назва", "") or r.get("Школа", "") or r.get("Покупна назва", "")).strip()
                    ph = str(r.get("Телефон", "") or r.get("Телефони", "")).strip()
                    add_record(em, ed, nm, ph, f"Ref Output ({fname})")
            except Exception:
                pass

    logger.info(
        "Master Email Lookups Ready: {} EDRPOU keys, {} Company Name keys, {} Phone keys!",
        len(edrpou_map), len(name_map), len(phone_map)
    )
    return edrpou_map, name_map, phone_map


def format_excel(ws, df: pd.DataFrame, header_color: str = "1F4E78"):
    ws.freeze_panes = "A2"
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    for col_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, col_name in enumerate(df.columns, 1):
        sample_vals = df[col_name].dropna().astype(str).head(100)
        max_len = max([len(str(v)) for v in sample_vals] + [len(col_name)])
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)


def enrich_1c_file(
    input_file: Path,
    output_excel: Path,
    output_csv: Path,
    edrpou_map: dict,
    name_map: dict,
    phone_map: dict,
    db_name: str,
    header_color: str
):
    if not input_file.exists():
        logger.warning("Input file {} not found!", input_file)
        return

    logger.info("Enriching 1C Database [{}] from {}...", db_name, input_file.name)
    df = pd.read_excel(input_file)

    # Detect key columns
    edr_col = None
    name_col = None
    email_col = None
    phone_col = None

    for c in df.columns:
        cl = c.lower()
        if "єдрпоу" in cl or "код" in cl:
            edr_col = c
        elif "контрагент" in cl or "назва" in cl:
            name_col = c
        elif "email" in cl or "пошта" in cl:
            email_col = c
        elif "телефон" in cl:
            phone_col = c

    if not name_col and len(df.columns) >= 2:
        name_col = df.columns[1]

    enriched_emails = []
    enrichment_sources = []
    match_types = []

    matched_by_edrpou = 0
    matched_by_name = 0
    matched_by_phone = 0
    already_valid = 0

    for _, r in df.iterrows():
        existing_email = str(r.get(email_col, "")).strip() if email_col else ""
        raw_edr = str(r.get(edr_col, "")).strip() if edr_col else ""
        raw_name = str(r.get(name_col, "")).strip() if name_col else ""
        raw_phone = str(r.get(phone_col, "")).strip() if phone_col else ""

        # 1. Existing valid email
        if is_valid_email_format(existing_email):
            enriched_emails.append(existing_email)
            enrichment_sources.append("1С Базовий")
            match_types.append("Існуючий в 1С")
            already_valid += 1
            continue

        # 2. Match by EDRPOU
        edr = clean_edrpou(raw_edr)
        if edr and edr in edrpou_map:
            em, src = edrpou_map[edr]
            enriched_emails.append(em)
            enrichment_sources.append(f"Збагачено з {src}")
            match_types.append("За ЄДРПОУ")
            matched_by_edrpou += 1
            continue

        # 3. Match by Company Name
        cn = clean_company_name(raw_name)
        if cn and cn in name_map:
            em, src = name_map[cn]
            enriched_emails.append(em)
            enrichment_sources.append(f"Збагачено з {src}")
            match_types.append("За Назвою")
            matched_by_name += 1
            continue

        # 4. Match by Phone
        p9 = clean_phone_last9(raw_phone)
        if p9 and p9 in phone_map:
            em, src = phone_map[p9]
            enriched_emails.append(em)
            enrichment_sources.append(f"Збагачено з {src}")
            match_types.append("За телефоном")
            matched_by_phone += 1
            continue

        enriched_emails.append("")
        enrichment_sources.append("")
        match_types.append("")

    df["Email"] = enriched_emails
    df["Джерело Email"] = enrichment_sources
    df["Тип Метчингу"] = match_types

    total_with_email = sum(1 for e in enriched_emails if e)
    newly_enriched = total_with_email - already_valid

    logger.info(
        "1C [{}] Enrichment Summary: Total rows={}, Valid original={}, Newly Enriched={} "
        "(By EDRPOU: {}, By Name: {}, By Phone: {}), Total Emails={}",
        db_name, len(df), already_valid, newly_enriched, matched_by_edrpou, matched_by_name, matched_by_phone, total_with_email
    )

    # Save to CSV and Excel
    df.to_csv(output_csv, index=False, encoding="utf-8-sig")

    with pd.ExcelWriter(output_excel, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=f"1С {db_name}")
        format_excel(writer.sheets[f"1С {db_name}"], df, header_color=header_color)

    logger.info("Saved enriched 1C [{}] dataset -> {} and {}", db_name, output_excel.name, output_csv.name)


def main():
    edrpou_map, name_map, phone_map = build_master_scraped_email_lookups()

    # 1. Enrich 1C UTP
    utp_in = OUTPUT_DIR / "1c_utp_full_contacts.xlsx"
    if not utp_in.exists():
        utp_in = OUTPUT_DIR / "1c_base_utp_ipt_contacts.xlsx"

    utp_out_excel = OUTPUT_DIR / "1c_utp_full_contacts_with_verified_emails.xlsx"
    utp_out_csv = OUTPUT_DIR / "1c_utp_full_contacts_with_verified_emails.csv"
    enrich_1c_file(utp_in, utp_out_excel, utp_out_csv, edrpou_map, name_map, phone_map, "UTP", "1F4E78")

    # 2. Enrich 1C Magazin
    mag_in = OUTPUT_DIR / "1c_magazin_full_contacts.xlsx"
    if not mag_in.exists():
        mag_in = OUTPUT_DIR / "1c_base_magazin_contacts.xlsx"

    mag_out_excel = OUTPUT_DIR / "1c_magazin_full_contacts_with_verified_emails.xlsx"
    mag_out_csv = OUTPUT_DIR / "1c_magazin_full_contacts_with_verified_emails.csv"
    enrich_1c_file(mag_in, mag_out_excel, mag_out_csv, edrpou_map, name_map, phone_map, "Magazin", "2E75B6")


if __name__ == "__main__":
    main()
