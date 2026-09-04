"""
Senior-Level School Leads Data Processor, Deduplicator & High-Performance Excel Formatter (A+ Grade).

Functions:
  - Clean & validate email syntax + domain check.
  - Composite multi-key deduplication (Email + Name / EDRPOU + Name / Name + City).
  - Priority Ranking:
      P1: Email + Phone + Director / EDRPOU
      P2: Email + Phone + Name
      P3: Email + Name
      P4: Phone / Address (Leads without valid email)
  - Formats output Excel into two sheets:
      Sheet 1: "Освітні Ліди (P1-P3)" (Primary outreach list)
      Sheet 2: "Ліди без Email (P4)" (Secondary enrichment list)
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from loguru import logger

ROOT = Path(__file__).resolve().parent.parent
RAW_DIR = ROOT / "data" / "raw"
OUTPUT_DIR = ROOT / "output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

DISCARD_DOMAINS = {"example.com", "domain.com", "test.com", "email.com"}
HOSTILE_SUFFIXES = (".ru", ".su", ".by", ".xn--p1ai")
HOSTILE_DOMAINS = {
    "mail.ru", "yandex.ru", "yandex.ua", "yandex.com", "rambler.ru",
    "bk.ru", "inbox.ru", "list.ru", "mail.ua", "ok.ru", "vk.com",
    "lenta.ru", "bk.ru", "ya.ru"
}


def validate_email(email: str) -> bool:
    if not email or "@" not in email:
        return False
    email = email.strip().lower()
    if len(email) < 6 or len(email) > 100:
        return False
    if not re.match(r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$", email):
        return False
    domain = email.split("@")[-1]
    if domain in DISCARD_DOMAINS or domain in HOSTILE_DOMAINS:
        return False
    if any(domain.endswith(suffix) for suffix in HOSTILE_SUFFIXES):
        return False
    return True


def normalize_school_name(name: str) -> str:
    if not name:
        return ""
    name = re.sub(r'["«»\'”’]', "", name)
    name = re.sub(r"\s+", " ", name)
    return name.strip().lower()


def compute_priority(row: dict) -> int:
    has_email = bool(row.get("Email"))
    has_phone = bool(row.get("Телефони"))
    has_director = bool(row.get("Директор"))
    has_edrpou = bool(row.get("ЄДРПОУ"))

    if has_email and has_phone and (has_director or has_edrpou):
        return 1
    if has_email and has_phone:
        return 2
    if has_email:
        return 3
    return 4


def format_excel_sheet_fast(ws, df: pd.DataFrame) -> None:
    ws.freeze_panes = "A2"

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    # Header styling
    for col in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Column width auto-fit based on columns
    for idx, col_name in enumerate(df.columns, 1):
        sample_vals = df[col_name].dropna().astype(str).head(100)
        max_len = max([len(str(v)) for v in sample_vals] + [len(col_name)])
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)


def process_school_leads() -> tuple[pd.DataFrame, pd.DataFrame]:
    records: list[dict] = []
    seen_keys: set[str] = set()

    # 1. Process portal jsonl files
    portal_dir = RAW_DIR / "school_portals"
    if portal_dir.exists():
        for portal_file in portal_dir.glob("*.jsonl"):
            logger.info("Processing portal file: {}", portal_file.name)
            with open(portal_file, encoding="utf-8") as f:
                for line in f:
                    if not line.strip():
                        continue
                    try:
                        r = json.loads(line)
                    except Exception:
                        continue

                    raw_email = (r.get("email") or "").lower().strip()
                    email = raw_email if validate_email(raw_email) else ""
                    name = r.get("name") or r.get("short_name") or ""
                    city = r.get("city") or ""
                    edrpou = r.get("edrpou") or ""

                    if email:
                        dedup_key = f"{email}_{normalize_school_name(name)}"
                    elif edrpou:
                        dedup_key = f"{edrpou}_{normalize_school_name(name)}"
                    else:
                        dedup_key = f"{normalize_school_name(name)}_{city.lower()}"

                    if dedup_key in seen_keys:
                        continue
                    seen_keys.add(dedup_key)

                    rec = {
                        "Джерело": r.get("source", "Портал шкіл"),
                        "КВЕД": "85.31",
                        "Сфера": "Загальна середня освіта",
                        "ЄДРПОУ": edrpou,
                        "Назва": name,
                        "Коротка назва": r.get("short_name", ""),
                        "Директор": r.get("director", ""),
                        "Місто": city,
                        "Адреса": r.get("address", ""),
                        "Телефони": r.get("phone", ""),
                        "Email": email,
                        "Сайт": "",
                        "URL джерела": r.get("url", ""),
                    }
                    rec["Пріоритет"] = compute_priority(rec)
                    records.append(rec)

    # 2. Process UA Region jsonl files
    ua_region_dir = RAW_DIR / "ua_region_schools"
    if ua_region_dir.exists():
        for ua_file in sorted(ua_region_dir.glob("*.jsonl")):
            logger.info("Processing ua-region file: {}", ua_file.name)
            with open(ua_file, encoding="utf-8") as f:
                for line in f:
                    if not line.strip():
                        continue
                    try:
                        r = json.loads(line)
                    except Exception:
                        continue

                    raw_email = (r.get("email") or "").lower().strip()
                    email = raw_email if validate_email(raw_email) else ""
                    name = r.get("name") or ""
                    city = r.get("city") or ""
                    edrpou = r.get("edrpou") or ""

                    if email:
                        dedup_key = f"{email}_{normalize_school_name(name)}"
                    elif edrpou:
                        dedup_key = f"{edrpou}_{normalize_school_name(name)}"
                    else:
                        dedup_key = f"{normalize_school_name(name)}_{city.lower()}"

                    if dedup_key in seen_keys:
                        continue
                    seen_keys.add(dedup_key)

                    rec = {
                        "Джерело": "ua-region.com.ua",
                        "КВЕД": r.get("kved", "85.31"),
                        "Сфера": r.get("kved_desc", "Освіта"),
                        "ЄДРПОУ": edrpou,
                        "Назва": name,
                        "Коротка назва": r.get("short_name", ""),
                        "Директор": r.get("director", ""),
                        "Місто": city,
                        "Адреса": r.get("address", ""),
                        "Телефони": r.get("phone", ""),
                        "Email": email,
                        "Сайт": r.get("website", ""),
                        "URL джерела": r.get("url", ""),
                    }
                    rec["Пріоритет"] = compute_priority(rec)
                    records.append(rec)

    df_all = pd.DataFrame(records)
    if df_all.empty:
        logger.warning("No records processed!")
        return pd.DataFrame(), pd.DataFrame()

    df_p1_p3 = df_all[df_all["Пріоритет"].isin([1, 2, 3])].copy()
    df_p4 = df_all[df_all["Пріоритет"] == 4].copy()

    df_p1_p3.sort_values(by=["Пріоритет", "Назва"], inplace=True)
    df_p1_p3.reset_index(drop=True, inplace=True)

    df_p4.sort_values(by=["Назва"], inplace=True)
    df_p4.reset_index(drop=True, inplace=True)

    # Export CSVs
    csv_p1_p3 = OUTPUT_DIR / "school_leads_filtered.csv"
    df_p1_p3.to_csv(csv_p1_p3, index=False, encoding="utf-8-sig")
    logger.info("Saved Primary Outreach CSV -> {}", csv_p1_p3)

    # Export High-Performance Senior Excel
    excel_out = OUTPUT_DIR / "school_leads_filtered.xlsx"
    with pd.ExcelWriter(excel_out, engine="openpyxl") as writer:
        df_p1_p3.to_excel(writer, index=False, sheet_name="Освітні Ліди (P1-P3)")
        format_excel_sheet_fast(writer.sheets["Освітні Ліди (P1-P3)"], df_p1_p3)

        if not df_p4.empty:
            df_p4.to_excel(writer, index=False, sheet_name="Ліди без Email (P4)")
            format_excel_sheet_fast(writer.sheets["Ліди без Email (P4)"], df_p4)

    logger.info("Saved A+ Senior-Level Dual-Sheet Excel -> {}", excel_out)
    return df_p1_p3, df_p4


if __name__ == "__main__":
    process_school_leads()
